"""Exam scoring engine, result calculation, analytics, and publication service."""

import io
import csv
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any, Tuple
from uuid import UUID
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
from app.models.result import Result, ResultDetail, ResultPublication
from app.models.attempt import Attempt, AttemptQuestion, AttemptOption, Answer, AttemptState
from app.models.exam import Exam, ExamStatus
from app.models.enrollment import Enrollment
from app.models.question import QuestionVersion, QuestionOption
from app.models.stream import Stream
from app.schemas.result import PublicationConfigRequest


class ResultService:

    def score_attempt(self, db: Session, attempt_id: UUID) -> Result:
        """
        Deterministically calculate scores for a submitted attempt.
        Evaluates answers against QuestionVersion snapshots and saves stream breakdowns.
        """
        attempt = db.query(Attempt).filter(Attempt.id == attempt_id).first()
        if not attempt:
            raise ValueError("Attempt not found")

        exam = db.query(Exam).filter(Exam.id == attempt.exam_id).first()
        if not exam:
            raise ValueError("Exam not found")

        # Load attempt questions with answers, options, and version options
        attempt_questions = (
            db.query(AttemptQuestion)
            .options(
                joinedload(AttemptQuestion.answer),
                joinedload(AttemptQuestion.options),
                joinedload(AttemptQuestion.question_version).joinedload(QuestionVersion.options),
            )
            .filter(AttemptQuestion.attempt_id == attempt.id)
            .order_by(AttemptQuestion.sequence_number.asc())
            .all()
        )

        total_marks = 0.0
        total_obtained = 0.0
        correct_count = 0
        incorrect_count = 0
        unanswered_count = 0

        # Stream aggregation map: stream_id -> {total_marks, obtained_marks, correct, incorrect, unanswered}
        stream_map: Dict[UUID, Dict[str, Any]] = {}

        for aq in attempt_questions:
            st_id = aq.stream_id
            if st_id not in stream_map:
                stream_map[st_id] = {
                    "total_marks": 0.0,
                    "obtained_marks": 0.0,
                    "correct": 0,
                    "incorrect": 0,
                    "unanswered": 0,
                }

            q_marks = aq.marks or 1.0
            q_neg = aq.negative_marks or 0.0

            total_marks += q_marks
            stream_map[st_id]["total_marks"] += q_marks

            sel_opt_id = aq.answer.selected_option_id if aq.answer else None
            ver = aq.question_version
            correct_opt = next((opt for opt in (ver.options if ver else []) if opt.is_correct), None)

            if not sel_opt_id:
                # Unanswered
                unanswered_count += 1
                stream_map[st_id]["unanswered"] += 1
            elif correct_opt and sel_opt_id == correct_opt.id:
                # Correct
                correct_count += 1
                total_obtained += q_marks
                stream_map[st_id]["correct"] += 1
                stream_map[st_id]["obtained_marks"] += q_marks
            else:
                # Incorrect
                incorrect_count += 1
                total_obtained -= q_neg
                stream_map[st_id]["incorrect"] += 1
                stream_map[st_id]["obtained_marks"] -= q_neg

        # Total questions
        total_questions = len(attempt_questions)
        percentage = max(0.0, round((total_obtained / total_marks * 100), 2)) if total_marks > 0 else 0.0

        # Pass/fail evaluation
        passed = None
        if exam.pass_percentage is not None:
            passed = percentage >= exam.pass_percentage

        # Upsert Result
        result = db.query(Result).filter(Result.attempt_id == attempt.id).first()
        if not result:
            result = Result(
                attempt_id=attempt.id,
                student_id=attempt.student_id,
                exam_id=exam.id,
            )
            db.add(result)
            db.flush()

        result.total_marks = total_marks
        result.total_obtained = total_obtained
        result.total_questions = total_questions
        result.correct_count = correct_count
        result.incorrect_count = incorrect_count
        result.unanswered_count = unanswered_count
        result.percentage = percentage
        result.passed = passed

        # Re-create ResultDetails for each stream
        db.query(ResultDetail).filter(ResultDetail.result_id == result.id).delete()
        db.flush()

        for s_id, s_stats in stream_map.items():
            rd = ResultDetail(
                result_id=result.id,
                stream_id=s_id,
                total_marks=s_stats["total_marks"],
                obtained_marks=s_stats["obtained_marks"],
                correct_count=s_stats["correct"],
                incorrect_count=s_stats["incorrect"],
                unanswered_count=s_stats["unanswered"],
            )
            db.add(rd)

        db.commit()
        db.refresh(result)
        return result

    def score_exam(self, db: Session, exam_id: UUID) -> int:
        """Batch score all submitted attempts for an exam."""
        attempts = (
            db.query(Attempt)
            .filter(
                Attempt.exam_id == exam_id,
                Attempt.state.in_([AttemptState.SUBMITTED, AttemptState.AUTO_SUBMITTED]),
            )
            .all()
        )
        scored_count = 0
        for att in attempts:
            self.score_attempt(db, att.id)
            scored_count += 1
        return scored_count

    def _score_unscored_attempts(self, db: Session, exam_id: UUID) -> int:
        """Score only submitted attempts that don't yet have a Result record."""
        existing_result_attempt_ids = (
            db.query(Result.attempt_id)
            .filter(Result.exam_id == exam_id)
        )
        unscored = (
            db.query(Attempt)
            .filter(
                Attempt.exam_id == exam_id,
                Attempt.state.in_([AttemptState.SUBMITTED, AttemptState.AUTO_SUBMITTED]),
                ~Attempt.id.in_(existing_result_attempt_ids),
            )
            .all()
        )
        for att in unscored:
            self.score_attempt(db, att.id)
        return len(unscored)

    def get_exam_results(self, db: Session, exam_id: UUID) -> List[Dict[str, Any]]:
        """List all candidate results for an exam with ranking and stream breakdowns."""
        # Only score attempts that don't yet have results (avoids re-scoring all on every view)
        self._score_unscored_attempts(db, exam_id)

        results = (
            db.query(Result)
            .options(
                joinedload(Result.attempt),
                joinedload(Result.details).joinedload(ResultDetail.stream),
            )
            .filter(Result.exam_id == exam_id)
            .order_by(Result.total_obtained.desc(), Result.percentage.desc())
            .all()
        )

        formatted: List[Dict[str, Any]] = []
        for rank_idx, r in enumerate(results, start=1):
            enr = (
                db.query(Enrollment)
                .filter(Enrollment.exam_id == exam_id, Enrollment.student_id == r.student_id)
                .first()
            )
            student_email = enr.student_email if enr else "Student"

            stream_details = []
            for dt in r.details:
                stream_details.append(
                    {
                        "stream_id": dt.stream_id,
                        "stream_code": dt.stream.code if dt.stream else "STREAM",
                        "stream_name": dt.stream.name if dt.stream else "Stream",
                        "total_marks": dt.total_marks,
                        "obtained_marks": dt.obtained_marks,
                        "correct_count": dt.correct_count,
                        "incorrect_count": dt.incorrect_count,
                        "unanswered_count": dt.unanswered_count,
                    }
                )

            formatted.append(
                {
                    "result_id": r.id,
                    "attempt_id": r.attempt_id,
                    "student_id": r.student_id,
                    "student_email": student_email,
                    "exam_id": r.exam_id,
                    "rank": rank_idx,
                    "total_marks": r.total_marks,
                    "total_obtained": r.total_obtained,
                    "percentage": r.percentage,
                    "passed": r.passed,
                    "total_questions": r.total_questions,
                    "correct_count": r.correct_count,
                    "incorrect_count": r.incorrect_count,
                    "unanswered_count": r.unanswered_count,
                    "stream_details": stream_details,
                }
            )

        return formatted

    def get_candidate_result(
        self, db: Session, exam_id: UUID, student_id: UUID, is_admin: bool = False
    ) -> Dict[str, Any]:
        """
        Fetch candidate's result.
        Enforces controlled publication gatekeeper for students.
        """
        exam = db.query(Exam).filter(Exam.id == exam_id).first()
        if not exam:
            raise ValueError("Exam not found")

        pub = self.get_or_create_publication(db, exam_id)

        # Gatekeeper: non-admins cannot access unreleased results
        if not is_admin and not pub.is_published:
            raise ValueError("Results for this examination have not been published yet.")

        # Find or calculate result
        result = db.query(Result).filter(Result.exam_id == exam_id, Result.student_id == student_id).first()
        if not result:
            # Check for submitted attempt
            attempt = (
                db.query(Attempt)
                .filter(
                    Attempt.exam_id == exam_id,
                    Attempt.student_id == student_id,
                    Attempt.state.in_([AttemptState.SUBMITTED, AttemptState.AUTO_SUBMITTED]),
                )
                .first()
            )
            if not attempt:
                raise ValueError("No submitted exam attempt found for this candidate.")
            result = self.score_attempt(db, attempt.id)

        # Calculate rank efficiently using count of higher-scoring results
        higher_count = (
            db.query(func.count(Result.id))
            .filter(
                Result.exam_id == exam_id,
                Result.total_obtained > result.total_obtained,
            )
            .scalar()
            or 0
        )
        rank = higher_count + 1

        # Load stream details
        details = (
            db.query(ResultDetail)
            .options(joinedload(ResultDetail.stream))
            .filter(ResultDetail.result_id == result.id)
            .all()
        )
        stream_details = [
            {
                "stream_id": dt.stream_id,
                "stream_code": dt.stream.code if dt.stream else "STREAM",
                "stream_name": dt.stream.name if dt.stream else "Stream",
                "total_marks": dt.total_marks,
                "obtained_marks": dt.obtained_marks,
                "correct_count": dt.correct_count,
                "incorrect_count": dt.incorrect_count,
                "unanswered_count": dt.unanswered_count,
            }
            for dt in details
        ]

        # Load question breakdown if permitted
        questions_review = None
        can_view_answers = is_admin or pub.show_answers

        if can_view_answers:
            attempt_questions = (
                db.query(AttemptQuestion)
                .options(
                    joinedload(AttemptQuestion.answer),
                    joinedload(AttemptQuestion.options),
                    joinedload(AttemptQuestion.question_version).joinedload(QuestionVersion.options),
                )
                .filter(AttemptQuestion.attempt_id == result.attempt_id)
                .order_by(AttemptQuestion.sequence_number.asc())
                .all()
            )

            questions_review = []
            for aq in attempt_questions:
                ver = aq.question_version
                stream_obj = db.query(Stream).filter(Stream.id == aq.stream_id).first()
                sel_opt_id = aq.answer.selected_option_id if aq.answer else None
                correct_opt = next((opt for opt in (ver.options if ver else []) if opt.is_correct), None)

                # Find labels & texts
                sel_opt = next((opt for opt in (ver.options if ver else []) if opt.id == sel_opt_id), None)
                sel_label = None
                sel_text = None
                if sel_opt:
                    # Find candidate display label from attempt_options
                    att_opt = next((ao for ao in aq.options if ao.original_option_id == sel_opt.id), None)
                    sel_label = att_opt.display_label if att_opt else sel_opt.option_label
                    sel_text = sel_opt.option_text

                corr_label = None
                corr_text = None
                if correct_opt:
                    corr_att = next((ao for ao in aq.options if ao.original_option_id == correct_opt.id), None)
                    corr_label = corr_att.display_label if corr_att else correct_opt.option_label
                    corr_text = correct_opt.option_text

                is_correct = bool(sel_opt_id and correct_opt and sel_opt_id == correct_opt.id)

                show_explanation = is_admin or pub.show_explanations
                questions_review.append(
                    {
                        "sequence_number": aq.sequence_number,
                        "question_text": ver.question_text if ver else "",
                        "stream_name": stream_obj.name if stream_obj else "Subject",
                        "marks": aq.marks,
                        "negative_marks": aq.negative_marks,
                        "selected_option_label": sel_label,
                        "selected_option_text": sel_text,
                        "correct_option_label": corr_label,
                        "correct_option_text": corr_text,
                        "is_correct": is_correct if sel_opt_id else False,
                        "explanation": ver.explanation if (show_explanation and ver) else None,
                    }
                )

        enr = (
            db.query(Enrollment)
            .filter(Enrollment.exam_id == exam_id, Enrollment.student_id == student_id)
            .first()
        )

        return {
            "result_id": result.id,
            "attempt_id": result.attempt_id,
            "student_id": result.student_id,
            "student_email": enr.student_email if enr else "Candidate",
            "exam_id": result.exam_id,
            "exam_title": exam.title,
            "rank": rank,
            "total_marks": result.total_marks,
            "total_obtained": result.total_obtained,
            "percentage": result.percentage,
            "passed": result.passed,
            "total_questions": result.total_questions,
            "correct_count": result.correct_count,
            "incorrect_count": result.incorrect_count,
            "unanswered_count": result.unanswered_count,
            "stream_details": stream_details,
            "questions": questions_review,
        }

    def get_exam_analytics(self, db: Session, exam_id: UUID) -> Dict[str, Any]:
        """Aggregate statistical metrics, class performance, and stream averages."""
        exam = db.query(Exam).filter(Exam.id == exam_id).first()
        if not exam:
            raise ValueError("Exam not found")

        self._score_unscored_attempts(db, exam_id)
        results = db.query(Result).filter(Result.exam_id == exam_id).all()

        total_candidates = len(results)
        if total_candidates == 0:
            return {
                "exam_id": exam.id,
                "exam_title": exam.title,
                "total_candidates": 0,
                "average_score": 0.0,
                "highest_score": 0.0,
                "lowest_score": 0.0,
                "average_percentage": 0.0,
                "pass_percentage": exam.pass_percentage,
                "pass_rate": 0.0,
                "stream_averages": [],
                "score_brackets": {
                    "0-20%": 0,
                    "21-40%": 0,
                    "41-60%": 0,
                    "61-80%": 0,
                    "81-100%": 0,
                },
            }

        scores = [r.total_obtained for r in results]
        percentages = [r.percentage for r in results]

        avg_score = round(sum(scores) / total_candidates, 2)
        high_score = round(max(scores), 2)
        low_score = round(min(scores), 2)
        avg_pct = round(sum(percentages) / total_candidates, 2)

        passed_candidates = sum(1 for r in results if r.passed is True)
        pass_rate = (
            round((passed_candidates / total_candidates * 100), 2)
            if exam.pass_percentage is not None
            else 100.0
        )

        # Brackets
        brackets = {
            "0-20%": sum(1 for p in percentages if p <= 20.0),
            "21-40%": sum(1 for p in percentages if 20.0 < p <= 40.0),
            "41-60%": sum(1 for p in percentages if 40.0 < p <= 60.0),
            "61-80%": sum(1 for p in percentages if 60.0 < p <= 80.0),
            "81-100%": sum(1 for p in percentages if p > 80.0),
        }

        # Stream averages
        stream_details = (
            db.query(ResultDetail)
            .options(joinedload(ResultDetail.stream))
            .join(Result, ResultDetail.result_id == Result.id)
            .filter(Result.exam_id == exam_id)
            .all()
        )

        st_map: Dict[UUID, Dict[str, Any]] = {}
        for dt in stream_details:
            sid = dt.stream_id
            if sid not in st_map:
                st_map[sid] = {
                    "stream_id": sid,
                    "stream_code": dt.stream.code if dt.stream else "STREAM",
                    "stream_name": dt.stream.name if dt.stream else "Stream",
                    "total_obtained": 0.0,
                    "max_marks": dt.total_marks,
                    "count": 0,
                }
            st_map[sid]["total_obtained"] += dt.obtained_marks
            st_map[sid]["count"] += 1

        stream_averages = []
        for s_data in st_map.values():
            c = s_data["count"] or 1
            avg_ob = round(s_data["total_obtained"] / c, 2)
            max_m = s_data["max_marks"] or 1.0
            avg_p = round((avg_ob / max_m * 100), 2) if max_m > 0 else 0.0
            stream_averages.append(
                {
                    "stream_id": s_data["stream_id"],
                    "stream_code": s_data["stream_code"],
                    "stream_name": s_data["stream_name"],
                    "average_score": avg_ob,
                    "max_marks": max_m,
                    "average_percentage": avg_p,
                }
            )

        return {
            "exam_id": exam.id,
            "exam_title": exam.title,
            "total_candidates": total_candidates,
            "average_score": avg_score,
            "highest_score": high_score,
            "lowest_score": low_score,
            "average_percentage": avg_pct,
            "pass_percentage": exam.pass_percentage,
            "pass_rate": pass_rate,
            "stream_averages": stream_averages,
            "score_brackets": brackets,
        }

    def get_or_create_publication(self, db: Session, exam_id: UUID) -> ResultPublication:
        """Fetch or initialize publication record for an exam."""
        pub = db.query(ResultPublication).filter(ResultPublication.exam_id == exam_id).first()
        if not pub:
            pub = ResultPublication(
                exam_id=exam_id,
                is_published=False,
                show_answers=False,
                show_explanations=False,
            )
            db.add(pub)
            db.commit()
            db.refresh(pub)
        return pub

    def update_publication(
        self, db: Session, exam_id: UUID, config: PublicationConfigRequest, published_by_id: UUID
    ) -> ResultPublication:
        """Update publication rules; scores all submitted attempts upon publication."""
        pub = self.get_or_create_publication(db, exam_id)
        now = datetime.now(timezone.utc)

        if config.is_published and not pub.is_published:
            pub.published_at = now
            pub.published_by_id = published_by_id
            # Ensure every submitted paper is evaluated
            self.score_exam(db, exam_id)

        pub.is_published = config.is_published
        pub.show_answers = config.show_answers
        pub.show_explanations = config.show_explanations

        db.commit()
        db.refresh(pub)
        return pub

    def _get_exam_streams_metadata(self, db: Session, exam_id: UUID, results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Collect ordered unique streams with names, codes, and max marks for this exam."""
        from app.models.exam import ExamStream
        streams_list = []
        seen_ids = set()

        # First collect streams from student result details (has exact total_marks calculated)
        for r in results:
            for sd in r.get("stream_details", []):
                sid_str = str(sd.get("stream_id"))
                if sid_str and sid_str not in seen_ids:
                    seen_ids.add(sid_str)
                    streams_list.append({
                        "stream_id": sid_str,
                        "stream_name": sd.get("stream_name") or "Subject Stream",
                        "stream_code": sd.get("stream_code") or "STREAM",
                        "total_marks": sd.get("total_marks") or 0.0,
                    })

        # Supplement from configured exam streams if not already present
        es_rows = (
            db.query(ExamStream)
            .options(joinedload(ExamStream.stream))
            .filter(ExamStream.exam_id == exam_id)
            .order_by(ExamStream.created_at.asc())
            .all()
        )
        for es in es_rows:
            sid_str = str(es.stream_id)
            if sid_str not in seen_ids:
                seen_ids.add(sid_str)
                calc_marks = float((es.question_count or 0) * (es.marks_per_question or 1.0))
                streams_list.append({
                    "stream_id": sid_str,
                    "stream_name": es.stream.name if es.stream else "Subject Stream",
                    "stream_code": es.stream.code if es.stream else "STREAM",
                    "total_marks": calc_marks,
                })

        return streams_list

    def generate_csv_export(self, db: Session, exam_id: UUID) -> str:
        """Export ranked candidate results to CSV format with dedicated stream-wise columns (Format A)."""
        results = self.get_exam_results(db, exam_id)
        exam = db.query(Exam).filter(Exam.id == exam_id).first()
        exam_title = exam.title if exam else "Exam"

        streams_list = self._get_exam_streams_metadata(db, exam_id, results)

        output = io.StringIO()
        writer = csv.writer(output)

        # Header with dedicated stream columns
        header = [
            "Rank",
            "Student Email",
            "Exam Title",
            "Score Obtained",
            "Total Marks",
            "Percentage",
            "Status",
        ]
        for s in streams_list:
            header.append(f"{s['stream_name']} ({s['stream_code']} - Max {s['total_marks']})")

        header.extend(["Correct", "Incorrect", "Unanswered"])
        writer.writerow(header)

        for r in results:
            status_str = "Passed" if r["passed"] is True else "Failed" if r["passed"] is False else "Completed"
            st_score_map = {str(sd["stream_id"]): sd["obtained_marks"] for sd in r.get("stream_details", [])}

            row = [
                r["rank"],
                r["student_email"],
                exam_title,
                r["total_obtained"],
                r["total_marks"],
                f"{r['percentage']}%",
                status_str,
            ]
            for s in streams_list:
                row.append(st_score_map.get(s["stream_id"], 0.0))

            row.extend([
                r["correct_count"],
                r["incorrect_count"],
                r["unanswered_count"],
            ])
            writer.writerow(row)

        return output.getvalue()

    def generate_excel_export(self, db: Session, exam_id: UUID) -> bytes:
        """Export ranked candidate results to professionally styled Excel (.xlsx) format with dedicated stream columns (Format A)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        results = self.get_exam_results(db, exam_id)
        exam = db.query(Exam).filter(Exam.id == exam_id).first()
        exam_title = exam.title if exam else "Exam"

        streams_list = self._get_exam_streams_metadata(db, exam_id, results)

        wb = Workbook()
        ws = wb.active
        ws.title = "Exam Results"

        # Theme styling
        header_fill = PatternFill(start_color="1F3B45", end_color="1F3B45", fill_type="solid")
        stream_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        headers = [
            "Rank",
            "Student Email",
            "Exam Title",
            "Score Obtained",
            "Total Marks",
            "Percentage",
            "Status",
        ]
        stream_col_start = len(headers) + 1
        for s in streams_list:
            headers.append(f"{s['stream_name']} ({s['stream_code']} - Max {s['total_marks']})")
        stream_col_end = len(headers)

        headers.extend(["Correct", "Incorrect", "Unanswered"])
        ws.append(headers)

        # Style Header Row
        for col_idx, cell in enumerate(ws[1], start=1):
            cell.font = header_font
            if stream_col_start <= col_idx <= stream_col_end:
                cell.fill = stream_header_fill
            else:
                cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data Rows
        for r in results:
            status_str = "Passed" if r["passed"] is True else "Failed" if r["passed"] is False else "Completed"
            st_score_map = {str(sd["stream_id"]): sd["obtained_marks"] for sd in r.get("stream_details", [])}

            row_data = [
                r["rank"],
                r["student_email"],
                exam_title,
                r["total_obtained"],
                r["total_marks"],
                f"{r['percentage']}%",
                status_str,
            ]
            for s in streams_list:
                row_data.append(st_score_map.get(s["stream_id"], 0.0))

            row_data.extend([
                r["correct_count"],
                r["incorrect_count"],
                r["unanswered_count"],
            ])
            ws.append(row_data)

            row_idx = ws.max_row
            for col_idx, cell in enumerate(ws[row_idx], start=1):
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in [1, 6, 7]:
                    cell.alignment = center_align
                elif col_idx in [2, 3]:
                    cell.alignment = left_align
                else:
                    cell.alignment = right_align

            # Highlight status cell
            status_cell = ws.cell(row=row_idx, column=7)
            if status_str == "Passed":
                status_cell.font = Font(name="Arial", size=10, bold=True, color="047857")
            elif status_str == "Failed":
                status_cell.font = Font(name="Arial", size=10, bold=True, color="B91C1C")

            ws.row_dimensions[row_idx].height = 20

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()


result_service = ResultService()
